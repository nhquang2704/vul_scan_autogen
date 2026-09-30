import os
from docx import Document
from datetime import datetime
import re
import unicodedata

def generate_docx_report(target_host, report_markdown_content, filename_prefix="vulnerability_report"):

    if not report_markdown_content or not report_markdown_content.strip():
        error_msg = "[DOCXGenerator] Lỗi: Không có nội dung báo cáo (rỗng hoặc None) để tạo file DOCX."
        return error_msg 

    reports_dir = "reports"
    if not os.path.exists(reports_dir):
        try:
            os.makedirs(reports_dir)
        except OSError as e:
            error_msg = f"[DOCXGenerator] Lỗi: Không thể tạo thư mục '{reports_dir}': {e}"
            print(error_msg)
            return error_msg
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_filename = f"{filename_prefix}_{timestamp}"
    filename = os.path.join(reports_dir, f"{base_filename}.docx")
    
    display_target = target_host if target_host else "Mục tiêu không xác định"

    try:
        doc = Document()
        title_paragraph = doc.add_heading(level=0)
        run_title = title_paragraph.add_run(f'Báo cáo Đánh giá Lỗ hổng Bảo mật')
        run_title.bold = True
        doc.add_paragraph(f"Mục tiêu: {display_target}")
        doc.add_paragraph(f"Báo cáo được tạo vào: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph() 

        cleaned_report_content = "".join(
            ch for ch in report_markdown_content
            if unicodedata.category(ch)[0] != "C" or ch in ('\n', '\r', '\t')
        )

        current_paragraph_text = []
        code_block_active = False
        code_block_lines = []

        for line in cleaned_report_content.split('\n'):
            stripped_line = line.strip()

            if stripped_line.startswith("```"):
                if code_block_active: 
                    if current_paragraph_text:
                        doc.add_paragraph(" ".join(current_paragraph_text).strip())
                        current_paragraph_text = []
                    if code_block_lines:
                        p = doc.add_paragraph()
                        for code_line in code_block_lines:
                            run = p.add_run(code_line + '\n')
                            run.font.name = 'Courier New'
                    code_block_active = False
                    code_block_lines = []
                else: 
                    if current_paragraph_text:
                        doc.add_paragraph(" ".join(current_paragraph_text).strip())
                        current_paragraph_text = []
                    code_block_active = True
                continue

            if code_block_active:
                code_block_lines.append(line)
                continue

            heading_level = 0
            if stripped_line.startswith('# '): heading_level = 1
            elif stripped_line.startswith('## '): heading_level = 2
            elif stripped_line.startswith('### '): heading_level = 3
            elif stripped_line.startswith('#### '): heading_level = 4

            if heading_level > 0:
                if current_paragraph_text:
                    doc.add_paragraph(" ".join(current_paragraph_text).strip())
                    current_paragraph_text = []
                heading_text = re.sub(r'^#+\s*', '', stripped_line).strip()
                doc.add_heading(heading_text, level=heading_level)
            elif stripped_line.startswith('* ') or stripped_line.startswith('- '):
                if current_paragraph_text:
                    doc.add_paragraph(" ".join(current_paragraph_text).strip())
                    current_paragraph_text = []
                item_text = re.sub(r'^[\*\-]\s*', '', stripped_line).strip()
                doc.add_paragraph(item_text, style='ListBullet')
            elif re.fullmatch(r'^(\*\*\*|---|___)\s*$', stripped_line): 
                if current_paragraph_text:
                    doc.add_paragraph(" ".join(current_paragraph_text).strip())
                    current_paragraph_text = []
                doc.add_paragraph("---")
            elif not stripped_line: 
                if current_paragraph_text:
                    doc.add_paragraph(" ".join(current_paragraph_text).strip())
                    current_paragraph_text = []
            else:
                current_paragraph_text.append(line)
        
        if current_paragraph_text:
            doc.add_paragraph(" ".join(current_paragraph_text).strip())
        if code_block_active and code_block_lines:
            p = doc.add_paragraph()
            for code_line in code_block_lines:
                run = p.add_run(code_line + '\n')
                run.font.name = 'Courier New'

        doc.save(filename)
        success_msg = f"[DOCXGenerator] Báo cáo đã được lưu thành công vào: {filename}"
        print(success_msg) # Giữ lại thông báo này
        return success_msg

    except Exception as e:
        error_msg = f"[DOCXGenerator] Lỗi khi tạo/lưu file DOCX '{filename}': {e}"
        print(error_msg)
        return error_msg